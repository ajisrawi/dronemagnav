"""Write the PCBWay BOM as .xlsx from the .csv produced by make_pcbway.py
(KiCad's bundled Python has no openpyxl, so this runs with the system Python)."""
import csv
import io
import os

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "hardware", "dronemagnav", "fab", "dronemagnav-bom-pcbway.csv")
DST = SRC.replace(".csv", ".xlsx")

rows = list(csv.reader(io.open(SRC, encoding="utf-8-sig", newline="")))
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "BOM"
for i, r in enumerate(rows):
    if i:
        r = [int(r[0]), r[1], int(r[2])] + r[3:]
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True)
for i, wdt in enumerate([7, 34, 5, 22, 26, 48, 44, 6, 40], 1):
    ws.column_dimensions[get_column_letter(i)].width = wdt
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)
ws.freeze_panes = "A2"
wb.save(DST)
print("wrote", DST, f"({len(rows) - 1} lines)")
