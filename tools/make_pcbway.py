"""PCBWay assembly files for the DroneMagNav board.

Run with KiCad's Python (needs pcbnew):
  & "C:\\Program Files\\KiCad\\10.0\\bin\\python.exe" tools/make_pcbway.py

Writes into hardware/dronemagnav/fab/:
  dronemagnav-centroid-pcbway.csv   PCBWay centroid / pick-and-place layout:
      Designator, Footprint, Mid X, Mid Y, Ref X, Ref Y, Pad X, Pad Y,
      Layer (T/B), Rotation, Comment - mm, origin = board bottom-left corner,
      Y positive upwards, Pad X/Y = centre of pad 1 (PCBWay uses it to check
      orientation)
  dronemagnav-bom-pcbway.csv / .xlsx  PCBWay BOM template columns:
      Item #, Designator, Qty, Manufacturer, Mfg Part #, Description / Value,
      Package/Footprint, Type, Your Instructions / Notes
"""
import csv
import io
import os
import sys

import pcbnew

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HW = os.path.join(ROOT, "hardware", "dronemagnav")
PCB = os.path.join(HW, "dronemagnav.kicad_pcb")
FAB = os.path.join(HW, "fab")
BOM = os.path.join(ROOT, "docs", "BOM.csv")
OUT_CENTROID = os.path.join(FAB, "dronemagnav-centroid-pcbway.csv")
OUT_BOM = os.path.join(FAB, "dronemagnav-bom-pcbway")


def mm(v):
    return pcbnew.ToMM(v)


def ref_key(ref):
    letters = ref.rstrip("0123456789")
    digits = "".join(c for c in ref if c.isdigit())
    return (letters, int(digits) if digits else 0)


def main():
    board = pcbnew.LoadBoard(PCB)
    bb = board.GetBoardEdgesBoundingBox()
    x0, y1 = mm(bb.GetLeft()), mm(bb.GetBottom())      # bottom-left corner

    rows = []
    for fp in board.Footprints():
        if fp.GetAttributes() & pcbnew.FP_EXCLUDE_FROM_POS_FILES:
            continue
        ref = fp.GetReference()
        if ref.startswith("H") and not fp.Pads() or ref.startswith("#"):
            continue
        pos = fp.GetPosition()
        pads = list(fp.Pads())
        first = next((p for p in pads if p.GetNumber() in ("1", "A1", "A1B12")), pads[0] if pads else None)
        ppos = first.GetPosition() if first else pos
        side = "T" if fp.GetLayer() == pcbnew.F_Cu else "B"
        rows.append([ref, fp.GetFPID().GetLibItemName().wx_str(),
                     f"{mm(pos.x) - x0:.3f}", f"{y1 - mm(pos.y):.3f}",
                     f"{mm(pos.x) - x0:.3f}", f"{y1 - mm(pos.y):.3f}",
                     f"{mm(ppos.x) - x0:.3f}", f"{y1 - mm(ppos.y):.3f}",
                     side, f"{fp.GetOrientationDegrees():.0f}", fp.GetValue()])
    rows.sort(key=lambda r: ref_key(r[0]))
    with io.open(OUT_CENTROID, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Designator", "Footprint", "Mid X", "Mid Y", "Ref X", "Ref Y",
                    "Pad X", "Pad Y", "Layer", "Rotation", "Comment"])
        w.writerows(rows)
    print(f"wrote {OUT_CENTROID} ({len(rows)} parts, mm, origin bottom-left)")
    placed = {r[0] for r in rows}

    # ------------------------------------------------------------------ BOM
    bom = []
    with io.open(BOM, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            refs = [x.strip() for x in r["References"].split(",") if x.strip()]
            refs = [x for x in refs if x in placed]
            if not refs:
                continue
            refs.sort(key=ref_key)
            desc = r["Description"].strip()
            val = r["Value"].strip()
            if any(ch.isdigit() for ch in val) and val not in desc:
                desc = f"{val} - {desc}" if desc else val
            notes = r.get("Notes", "").strip()
            bom.append([len(bom) + 1, ", ".join(refs), len(refs), r["Manufacturer"].strip(),
                        r["MPN"].strip(), desc, r["Footprint"].strip(), "SMD", notes])
    header = ["Item #", "Designator", "Qty", "Manufacturer", "Mfg Part #",
              "Description / Value", "Package/Footprint", "Type", "Your Instructions / Notes"]
    with io.open(OUT_BOM + ".csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(bom)
    total = sum(r[2] for r in bom)
    print(f"wrote {OUT_BOM}.csv ({len(bom)} lines, {total} parts)")
    in_bom = {x for r in bom for x in r[1].split(", ")}
    missing = sorted(placed - in_bom, key=ref_key)
    if missing:
        print("WARNING placed parts not in BOM:", missing)
    assert total == len(rows), (total, len(rows))

    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "BOM"
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in bom:
            ws.append(r)
        widths = [7, 34, 5, 22, 26, 48, 44, 6, 40]
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        wb.save(OUT_BOM + ".xlsx")
        print(f"wrote {OUT_BOM}.xlsx")
    except ImportError:
        print("openpyxl not available: xlsx skipped")


if __name__ == "__main__":
    sys.exit(main())
